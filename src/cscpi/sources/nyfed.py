"""New York Fed Global Supply Chain Pressure Index (benchmark only)."""

from __future__ import annotations

import io

import polars as pl
import requests

GSCPI_URL = "https://www.newyorkfed.org/medialibrary/research/interactives/gscpi/downloads/gscpi_data.xlsx"


def fetch(url: str = GSCPI_URL) -> pl.DataFrame:
    """Download the GSCPI workbook and return month, value. Needs openpyxl."""
    from openpyxl import load_workbook

    r = requests.get(url, timeout=60, headers={"User-Agent": "cscpi/0.1"})
    r.raise_for_status()
    wb = load_workbook(io.BytesIO(r.content), data_only=True, read_only=True)
    rows = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            if len(row) >= 2 and hasattr(row[0], "year") and isinstance(row[1], (int, float)):
                rows.append(
                    (
                        row[0].date().replace(day=1) if hasattr(row[0], "date") else row[0].replace(day=1),
                        float(row[1]),
                    )
                )
        if rows:
            break
    if not rows:
        raise RuntimeError("GSCPI workbook layout changed: no date/value rows found")
    return pl.DataFrame(rows, schema=["month", "value"], orient="row").unique("month").sort("month")
